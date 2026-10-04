"""Tạo results.xlsx đúng schema GUIDE 6.1 từ log và prediction thật."""
from __future__ import annotations

import json
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


CODE_DIR = Path(__file__).resolve().parent
STARTER_DIR = CODE_DIR.parent
REPO_DIR = STARTER_DIR.parent
sys.path.insert(0, str(REPO_DIR))
from eval import compute_metrics, load_group, load_names, read_pred


SEEDS = (0, 1, 2)
TRAINING_INFO = {
    "T00": ("Baseline", "Công thức nền"),
    "T01": ("B-Augmentation", "basic → color jitter"),
    "T02": ("C-Loss", "CE → label smoothing ε=0.1"),
    "T03": ("D-Sampler", "shuffle → balanced sampler"),
    "T04": ("E-LR", "LR head 1e-3 → 1e-4"),
    "T05": ("F-EMA", "không EMA → EMA decay=0.999"),
    "T10": ("Combined", "T01 + T02 + T03 + T04 + T05"),
}


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def prediction_metrics(path: Path):
    pred = read_pred(str(path))
    return compute_metrics(pred.y_true, pred.y_pred, pred.probs)


def build_backbones() -> pd.DataFrame:
    source = pd.read_csv(STARTER_DIR / "backbone_results.csv")
    latency = pd.read_csv(STARTER_DIR / "backbone_latency.csv").set_index("exp_id")
    rows = []
    for _, row in source.sort_values("exp_id").iterrows():
        exp_id = row["exp_id"]
        cfg = read_json(STARTER_DIR / "runs" / exp_id / "seed0" / "config.json")
        bench = latency.loc[exp_id]
        rows.append({
            "exp_id": exp_id,
            "backbone": row["backbone"],
            "weight_tag": row["tag trọng số"],
            "params_m": row["#params (M)"],
            "gmac": row["GMAC"],
            "resolution": cfg["img_size"],
            "epochs": cfg["epochs"],
            "best_epoch": row["best_epoch"],
            "seed": row["seed"],
            "macro_f1_val": row["macro-F1 val"],
            "top1_val": row["top-1 val"],
            "seconds_per_epoch": row["giây/epoch"],
            "latency_batch1_p50_ms": bench["p50"],
            "latency_batch1_p95_ms": bench["p95"],
            "notes": "FP32; latency chỉ forward, warmup=10, iterations=100",
        })
    return pd.DataFrame(rows)


def build_training() -> pd.DataFrame:
    summaries = {}
    configs = {}
    for exp_id in TRAINING_INFO:
        exp_dir = STARTER_DIR / "runs" / exp_id
        if not exp_dir.is_dir():
            continue
        for seed_dir in sorted(exp_dir.glob("seed*")):
            seed = int(seed_dir.name.removeprefix("seed"))
            summaries[(exp_id, seed)] = read_json(seed_dir / "summary.json")
            configs[(exp_id, seed)] = read_json(seed_dir / "config.json")

    baseline = {seed: summaries[("T00", seed)]["val_macro_f1"] for seed in SEEDS}
    seed_summary = pd.read_csv(STARTER_DIR / "training_seed_summary.csv").set_index("exp_id")
    rows = []
    for (exp_id, seed), summary in sorted(summaries.items()):
        cfg = configs[(exp_id, seed)]
        axis, difference = TRAINING_INFO[exp_id]
        val_path = STARTER_DIR / "predictions" / f"{exp_id}_seed{seed}_val.csv"
        metrics = prediction_metrics(val_path)
        comparison = baseline.get(seed, baseline[0])
        aggregate = seed_summary.loc[exp_id] if exp_id in seed_summary.index else None
        rows.append({
            "exp_id": exp_id,
            "backbone": cfg["backbone"],
            "axis": axis,
            "difference_vs_T00": difference,
            "seed": seed,
            "macro_f1_val": summary["val_macro_f1"],
            "top1_val": summary["val_top1"],
            "delta_vs_T00": summary["val_macro_f1"] - comparison,
            "f1_chinee_apple_val": metrics["f1"][0],
            "f1_snake_weed_val": metrics["f1"][7],
            "best_epoch": summary["best_epoch"],
            "seconds_per_epoch": summary["seconds_per_epoch"],
            "n_seeds": int(aggregate["n_seeds"]) if aggregate is not None else 1,
            "macro_f1_mean": aggregate["macro_f1_mean"] if aggregate is not None else summary["val_macro_f1"],
            "macro_f1_std": aggregate["macro_f1_std"] if aggregate is not None else np.nan,
            "notes": "Controlled ablation" if exp_id not in {"T00", "T10"} else (
                "Baseline" if exp_id == "T00" else "Kết hợp giảm so với T00; không chọn"
            ),
        })
    return pd.DataFrame(rows)


def build_inference() -> pd.DataFrame:
    frame = pd.read_csv(STARTER_DIR / "inference_results.csv")
    frame.insert(
        3, "model_checkpoint",
        ["T00/seed0+seed1+seed2" if exp == "I05" else "T00/seed0" for exp in frame["exp_id"]],
    )
    frame = frame.rename(columns={
        "K": "k_views_or_models",
        "p50_ms": "latency_p50_ms",
        "p95_ms": "latency_p95_ms",
        "p99_ms": "latency_p99_ms",
    })
    ordered = [
        "exp_id", "method", "model_checkpoint", "k_views_or_models",
        "macro_f1_val", "top1_val", "ece_val", "nll_val",
        "latency_p50_ms", "latency_p95_ms", "latency_p99_ms",
        "images_per_s", "relative_cost", "gpu", "dtype", "batch", "img_size",
        "warmup", "iterations", "includes_preprocessing",
    ]
    return frame[ordered].sort_values("exp_id").reset_index(drop=True)


def build_final_and_per_class():
    cfg = read_json(STARTER_DIR / "runs" / "T00" / "seed0" / "config.json")
    labels_dir = Path(cfg["labels_dir"])
    test_csv = labels_dir / "test_subset0.csv"
    class_names = load_names(str(labels_dir / "labels.csv"))
    configurations = {
        "T00": "ConvNeXt-Tiny | T00 | I00 single-view 224",
        "F01": "ConvNeXt-Tiny | T00 | I04 single-view 256",
    }
    rows = []
    per_class_rows = []
    for tag in ("T00", "F01"):
        group = load_group(
            str(STARTER_DIR / "predictions" / f"{tag}_seed*_test.csv"),
            str(test_csv), ref_what="test",
        )
        val_metrics = {
            seed: prediction_metrics(STARTER_DIR / "predictions" / f"{tag}_seed{seed}_val.csv")
            for seed in SEEDS
        }
        for prediction, metrics in zip(group.preds, group.metrics):
            seed = prediction.seed
            rows.append({
                "exp_id": tag,
                "row_type": "seed",
                "configuration": configurations[tag],
                "seed": seed,
                "macro_f1_val": val_metrics[seed]["macro_f1"],
                "macro_f1_test": metrics["macro_f1"],
                "top1_test": metrics["top1"],
                "ece_test": metrics["ece"],
                "balanced_acc_test": metrics["balanced_acc"],
                "nll_test": metrics["nll"],
                "macro_f1_test_std": np.nan,
                "top1_test_std": np.nan,
            })
        rows.append({
            "exp_id": tag,
            "row_type": "mean ± std",
            "configuration": configurations[tag],
            "seed": "all",
            "macro_f1_val": np.mean([val_metrics[seed]["macro_f1"] for seed in SEEDS]),
            "macro_f1_test": group.summary["macro_f1"][0],
            "top1_test": group.summary["top1"][0],
            "ece_test": group.summary["ece"][0],
            "balanced_acc_test": group.summary["balanced_acc"][0],
            "nll_test": group.summary["nll"][0],
            "macro_f1_test_std": group.summary["macro_f1"][1],
            "top1_test_std": group.summary["top1"][1],
        })
        precision_mean, precision_std = group.summary["precision"]
        recall_mean, recall_std = group.summary["recall"]
        f1_mean, f1_std = group.summary["f1"]
        support = group.metrics[0]["support"].astype(int)
        for class_id, class_name in enumerate(class_names):
            per_class_rows.append({
                "exp_id": tag,
                "class_id": class_id,
                "class": class_name,
                "support": int(support[class_id]),
                "precision_mean": precision_mean[class_id],
                "precision_std": precision_std[class_id],
                "recall_mean": recall_mean[class_id],
                "recall_std": recall_std[class_id],
                "f1_mean": f1_mean[class_id],
                "f1_std": f1_std[class_id],
            })
    return pd.DataFrame(rows), pd.DataFrame(per_class_rows)


def build_latency() -> pd.DataFrame:
    methods = pd.read_csv(STARTER_DIR / "inference_results.csv").set_index("exp_id")["method"]
    frame = pd.read_csv(STARTER_DIR / "latency_results.csv")
    frame.insert(1, "configuration", [methods.get(exp, "1-view 224, batch 32") for exp in frame["exp_id"]])
    frame.insert(6, "fused_bn", "no")
    frame = frame.rename(columns={"p50_ms": "p50", "p95_ms": "p95", "p99_ms": "p99"})
    return frame[[
        "exp_id", "configuration", "gpu", "dtype", "batch", "img_size", "fused_bn",
        "p50", "p95", "p99", "images_per_s", "relative_cost",
        "warmup", "iterations", "includes_preprocessing",
    ]]


def build_summary(backbones, training, inference):
    rows = []
    for _, row in backbones.iterrows():
        rows.append({
            "stage": "Backbone", "exp_id": row.exp_id, "configuration": row.backbone,
            "macro_f1_val": row.macro_f1_val, "top1_val": row.top1_val,
            "p95_ms": row.latency_batch1_p95_ms,
            "relative_cost": np.nan, "notes": f"{row.params_m:.2f}M params; {row.gmac:.2f} GMAC",
        })
    for _, row in training.iterrows():
        rows.append({
            "stage": "Training", "exp_id": f"{row.exp_id}/s{row.seed}",
            "configuration": row.difference_vs_T00,
            "macro_f1_val": row.macro_f1_val, "top1_val": row.top1_val,
            "p95_ms": np.nan, "relative_cost": np.nan, "notes": row.axis,
        })
    for _, row in inference.iterrows():
        rows.append({
            "stage": "Inference", "exp_id": row.exp_id, "configuration": row.method,
            "macro_f1_val": row.macro_f1_val, "top1_val": row.top1_val,
            "p95_ms": row.latency_p95_ms, "relative_cost": row.relative_cost,
            "notes": f"K={row.k_views_or_models}",
        })
    result = pd.DataFrame(rows).sort_values("macro_f1_val", ascending=False).head(10).reset_index(drop=True)
    result.insert(0, "rank", np.arange(1, len(result) + 1))
    return result


def style_and_save(sheets: dict[str, pd.DataFrame], path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
        workbook = writer.book
        navy, white = "17365D", "FFFFFF"
        thin_gray = Side(style="thin", color="D9E2F3")
        for name, frame in sheets.items():
            ws = workbook[name]
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 30
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(color=white, bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = Border(bottom=thin_gray)
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
            for col_idx, column in enumerate(frame.columns, 1):
                values = [str(column)] + ["" if pd.isna(v) else str(v) for v in frame[column].head(200)]
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max(map(len, values)) + 2, 11), 44)
                lowered = str(column).lower()
                if any(token in lowered for token in ("f1", "top1", "ece", "nll", "precision", "recall", "balanced", "delta")):
                    for cell in ws[get_column_letter(col_idx)][1:]:
                        cell.number_format = "0.0000"
                elif any(token in lowered for token in ("p50", "p95", "p99", "seconds", "gmac", "params", "latency")):
                    for cell in ws[get_column_letter(col_idx)][1:]:
                        cell.number_format = "0.00"
            if not frame.empty:
                ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"
                table = Table(displayName=f"Tbl{name}", ref=ref)
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showFirstColumn=False,
                    showLastColumn=False, showRowStripes=True, showColumnStripes=False,
                )
                ws.add_table(table)
                for col_idx, column in enumerate(frame.columns, 1):
                    if "macro_f1" in str(column).lower():
                        letter = get_column_letter(col_idx)
                        ws.conditional_formatting.add(
                            f"{letter}2:{letter}{len(frame) + 1}",
                            ColorScaleRule(
                                start_type="min", start_color="F8696B",
                                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                end_type="max", end_color="63BE7B",
                            ),
                        )
        ws = workbook["Summary"]
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Top configurations by validation macro-F1"
        chart.x_axis.title = "Macro-F1"
        chart.y_axis.title = "Experiment"
        chart.height, chart.width = 8, 14
        macro_col = list(sheets["Summary"].columns).index("macro_f1_val") + 1
        exp_col = list(sheets["Summary"].columns).index("exp_id") + 1
        chart.add_data(Reference(ws, min_col=macro_col, min_row=1, max_row=len(sheets["Summary"]) + 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=exp_col, min_row=2, max_row=len(sheets["Summary"]) + 1))
        ws.add_chart(chart, "J2")


def verify(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != list(sheets):
        raise AssertionError(workbook.sheetnames)
    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
    errors = [
        f"{ws.title}!{cell.coordinate}={cell.value}"
        for ws in workbook.worksheets for row in ws.iter_rows() for cell in row
        if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens)
    ]
    if errors:
        raise AssertionError(errors)
    required_columns = {
        "Backbones": {"resolution", "epochs", "latency_batch1_p50_ms", "notes"},
        "Training": {"axis", "difference_vs_T00", "delta_vs_T00", "f1_chinee_apple_val", "notes"},
        "Inference": {"model_checkpoint", "latency_p50_ms", "latency_p95_ms", "latency_p99_ms"},
        "Final": {"configuration", "macro_f1_val", "macro_f1_test", "macro_f1_test_std"},
        "PerClass": {"class", "support", "precision_mean", "recall_mean", "f1_mean"},
        "Latency": {"configuration", "gpu", "dtype", "batch", "fused_bn", "p50", "p95", "p99", "images_per_s"},
        "Summary": {"rank", "exp_id", "macro_f1_val", "p95_ms"},
    }
    for name, expected in required_columns.items():
        headers = {cell.value for cell in workbook[name][1]}
        if not expected.issubset(headers):
            raise AssertionError(f"{name} thiếu {sorted(expected - headers)}")
    workbook.close()


def main() -> None:
    backbones = build_backbones()
    training = build_training()
    inference = build_inference()
    final, per_class = build_final_and_per_class()
    latency = build_latency()
    summary = build_summary(backbones, training, inference)
    sheets = {
        "Backbones": backbones,
        "Training": training,
        "Inference": inference,
        "Final": final,
        "PerClass": per_class,
        "Latency": latency,
        "Summary": summary,
    }
    output = STARTER_DIR / "results.xlsx"
    style_and_save(sheets, output)
    verify(output, sheets)
    print(f"Workbook written and verified: {output}")
    for name, frame in sheets.items():
        print(f"{name}: {len(frame)} rows, {len(frame.columns)} columns")


if __name__ == "__main__":
    main()
